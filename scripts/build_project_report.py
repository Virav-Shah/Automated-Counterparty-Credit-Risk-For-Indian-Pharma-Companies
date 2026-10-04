"""Build the finance interview report from the project's reproducible results.

Requires python-docx. Run from any directory; input and output paths are resolved
against this repository. Narrative findings are analytical interpretations, not
claims that full audited statements or actual counterparty exposures were supplied.
The companion Markdown export is the editable content source for plain-text review.
"""
from pathlib import Path
import json
import statistics
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'reports'
DATA = json.loads((ROOT / 'outputs/dashboard_data.json').read_text())
ROWS = DATA['rows']
LATEST = {r['company_id']: r for r in ROWS if r['fiscal_year'] == 'FY25'}
DOC = Document()
MD = []


def paragraph(text, style=None):
    """Add a paragraph and mirror its content to the Markdown export."""
    p = DOC.add_paragraph(style=style)
    url = re.search(r'https://\S+$', text)
    if url:
        p.add_run(text[:url.start()])
        link = OxmlElement('w:hyperlink')
        link.set(qn('r:id'), p.part.relate_to(url.group(), RT.HYPERLINK, is_external=True))
        run = OxmlElement('w:r')
        props = OxmlElement('w:rPr')
        color = OxmlElement('w:color'); color.set(qn('w:val'), '000000'); props.append(color)
        underline = OxmlElement('w:u'); underline.set(qn('w:val'), 'single'); props.append(underline)
        run.append(props)
        label = OxmlElement('w:t'); label.text = 'Source'; run.append(label)
        link.append(run); p._p.append(link)
    else:
        p.add_run(text)
    MD.append(text + '\n')


def heading(text, level=2):
    """Add a semantic heading with a corresponding Markdown heading."""
    DOC.add_heading(text, level)
    MD.append('#' * (level + 1) + ' ' + text + '\n')


def page(title, subtitle=None):
    """Start one explicitly paginated report section."""
    if DOC.paragraphs:
        DOC.add_page_break()
    heading(title, 1)
    if subtitle:
        paragraph(subtitle)


def table(headers, values, widths):
    """Add a fixed-width, accessible comparison table with expanding rows."""
    t = DOC.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, width in zip(t.columns, widths):
        col.width = Inches(width)
    for cell, label, width in zip(t.rows[0].cells, headers, widths):
        cell.width = Inches(width)
        cell.text = str(label)
    for vals in values:
        row = t.add_row()
        for cell, val, width in zip(row.cells, vals, widths):
            cell.width = Inches(width)
            cell.text = str(val)
    for ri, row in enumerate(t.rows):
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if ri == 0:
            trpr.append(OxmlElement('w:tblHeader'))
        for ci, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            borders = OxmlElement('w:tcBorders')
            for side in ['top', 'left', 'bottom', 'right']:
                el = OxmlElement('w:' + side)
                for k, v in {'val': 'single', 'sz': '4', 'color': 'D9D9D9'}.items():
                    el.set(qn('w:' + k), v)
                borders.append(el)
            tcpr.append(borders)
            margins = OxmlElement('w:tcMar')
            for side in ['top', 'bottom', 'left', 'right']:
                el = OxmlElement('w:' + side)
                el.set(qn('w:w'), '85')
                el.set(qn('w:type'), 'dxa')
                margins.append(el)
            tcpr.append(margins)
            fill = OxmlElement('w:shd')
            fill.set(qn('w:fill'), '18324D' if ri == 0 else ('F1F4F7' if ri % 2 == 0 else 'FFFFFF'))
            tcpr.append(fill)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.line_spacing = 1.0
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if ci == 0 else WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.size = Pt(10)
                    run.bold = ri == 0
                    run.font.color.rgb = RGBColor.from_string('FFFFFF' if ri == 0 else '000000')
    DOC.add_paragraph().paragraph_format.space_after = Pt(0)
    MD.append('| ' + ' | '.join(headers) + ' |')
    MD.append('| ' + ' | '.join(['---'] * len(headers)) + ' |')
    MD.extend('| ' + ' | '.join(map(str, vals)) + ' |' for vals in values)
    MD.append('')


def number(value, digits=2):
    """Format an unavailable ratio distinctly from a numeric zero."""
    return 'n/a' if value is None else f'{value:,.{digits}f}'


def summary(name, comment):
    """Write a compact, issuer-specific financial assessment."""
    r = LATEST[name]
    text = f"FY25 score {r['financial_risk_score']}/100 ({r['risk_rating']}); debt/EBITDA {number(r['debt_ebitda'])}x; EBIT/interest {number(r['interest_coverage'])}x. " + comment
    p = DOC.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    p.add_run(name + '. ').bold = True
    p.add_run(text)
    MD.append('**' + name + '.** ' + text + '\n')


def case_table(name):
    """Compare the complete five-year trajectory of a detailed case study."""
    series = sorted([r for r in ROWS if r['company_id'] == name], key=lambda r: r['fiscal_year'])
    table(['Metric'] + [r['fiscal_year'] for r in series], [
        ['Revenue'] + [number(r['revenue'], 0) for r in series],
        ['EBITDA margin'] + [f"{r['ebitda']/r['revenue']:.1%}" for r in series],
        ['Debt / EBITDA'] + [number(r['debt_ebitda']) + 'x' for r in series],
        ['EBIT / interest'] + [number(r['interest_coverage']) + 'x' for r in series],
        ['Free cash flow'] + [number(r['free_cash_flow'], 0) for r in series],
        ['Current ratio'] + [number(r['current_ratio']) + 'x' for r in series],
        ['Financial score'] + [number(r['financial_risk_score'], 0) for r in series],
    ], [1.65, .92, .92, .92, .92, .92])
    paragraph('Amounts are in source units; currency and scale require reconciliation to issuer statements. FY21 has no full score because FY20 growth inputs are absent. Source: project financial extract and calculations [1].')


def build():
    """Author fifteen report pages and save DOCX plus a readable Markdown copy."""
    sec = DOC.sections[0]
    sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
    sec.top_margin = sec.bottom_margin = Inches(.65)
    sec.left_margin = sec.right_margin = Inches(.70)
    for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3']:
        st = DOC.styles[name]
        st.font.name = 'Calibri'
        st.font.color.rgb = RGBColor(0, 0, 0)
    normal = DOC.styles['Normal']
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.space_after = Pt(7)
    for name, size in [('Heading 1', 20), ('Heading 2', 12)]:
        st = DOC.styles[name]
        st.font.size = Pt(size)
        st.paragraph_format.space_before = Pt(10)
        st.paragraph_format.space_after = Pt(6)
    foot = sec.footer.paragraphs[0]
    foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    foot._p.append(field)
    DOC.core_properties.title = 'Counterparty Credit Analysis of Indian Pharmaceutical Companies'
    DOC.core_properties.subject = 'Financial analysis and counterparty credit review'
    DOC.core_properties.author = ''

    page('Counterparty Credit Analysis of Indian Pharmaceutical Companies', 'Professional analyst report | Financial years FY21 to FY25 | Review date 4 October 2026')
    heading('Executive assessment')
    paragraph('This study evaluates 20 Indian pharmaceutical and biopharma companies through five-year financial analysis, business-model peer benchmarking and internal credit-risk classification. The objective is to identify repayment capacity, emerging financial stress and the information needed before accepting or increasing counterparty exposure. The central analytical question is whether operating earnings convert into sufficient cash to support obligations through changing business conditions.')
    paragraph('The FY25 financial snapshot shows 15 companies in the VERY LOW internal band and five in LOW. The strongest measured pressure sits with Biocon at 25/100, followed by Piramal Pharma and Jubilant Pharmova at 22, Mankind at 20 and Laurus at 19. These classifications describe the historical extract; they do not establish present-day creditworthiness or authorize a limit. Historical stress was materially greater: Biocon reached 70 in FY23 and Laurus reached 70 in FY24. Their recoveries provide the two detailed case studies.')
    paragraph('External checks identified differences from Biocon and Laurus issuer disclosures, including a different sign for Laurus OCF minus capex. Results remain provisional pending source reconciliation. Pages 13-14 distinguish actual benchmark comparisons, passing calculation checks and unresolved source validation.')
    heading('What the analysis delivers')
    paragraph('The analysis connects P&L trends to leverage, interest servicing, liquidity, working-capital absorption and free cash flow. It compares each company against businesses with similar operating risks and produces financial early-warning flags. A separate market stress assessment supports review prioritization. Business and regulatory evidence remains narrative context; no qualitative score contributes to the rating or watchlist.')
    paragraph('Counterparty assessment extends beyond the listed company name to the actual obligor, contractual terms, group relationships and enforceability of protection. Credit-limit monitoring is specified as a practical framework. Actual approved limits, exposure balances and overdue data have not yet been incorporated, so no real utilization or breach result is presented.')
    heading('Reading guide')
    table(['Pages', 'Focus'], [['2-3', 'Sector economics and peer selection'], ['4-6', 'Financial analysis, rating methodology and portfolio findings'], ['7-8', 'Biocon and Laurus detailed case studies'], ['9-10', 'Summary assessments of the other 18 companies'], ['11-12', 'Credit reviews, counterparty assessment and credit-limit monitoring'], ['13-14', 'Benchmark comparisons, source reconciliation and validation'], ['15', 'Evidence standards, limitations and references']], [.8, 6.0])

    page('Sector analysis and credit transmission')
    paragraph('The Department of Pharmaceuticals reports pharmaceutical turnover of INR 471,898 crore in FY2024-25, annual growth of 13.07% and a 9.5% CAGR since FY21. Pharmaceutical exports were INR 245,962 crore and imports INR 63,573 crore in FY25. These national statistics describe the sector, while this project studies a selected 20-company sample [2, printed p. 3].')
    heading('Business economics')
    paragraph('Domestic branded formulations depend on therapy demand, distribution, prescription relationships and product portfolios. A broad brand base can support repeat sales, but pricing constraints, competitive spending and inventory across the channel still affect margins and collections. Credit analysis therefore tests whether sales growth is accompanied by stable earnings and operating cash generation.')
    paragraph('Global generics add geographic and product diversification, alongside competition, currency exposure and dependence on compliant manufacturing. Revenue growth can mask price pressure if launches offset erosion in mature products. The relevant credit questions concern earnings durability, geographic concentration, plant dependency and the cash cost of sustaining approvals and supply.')
    paragraph('API and contract development and manufacturing organizations (CDMO) face capacity utilization, customer programs and contract concentration. Investment can precede commercial cash inflows, making capex, debt funding and cash conversion central. Biologics, complex injectables and specialty businesses similarly require careful analysis of investment intensity, commercialization timelines and operational execution. These are analytical risk mechanisms, not claims that every studied issuer has the same exposure.')
    heading('How operating risk becomes credit risk')
    table(['Risk channel', 'Financial effect to investigate'], [['Pricing or demand pressure', 'Lower EBITDA and EBIT; weaker interest coverage'], ['Plant or regulatory disruption', 'Interrupted shipments, remediation cost, delayed receipts'], ['Customer concentration', 'Revenue volatility and correlated receivable exposure'], ['Expansion or acquisition', 'Higher debt, capex and refinancing requirements'], ['Inventory and collection pressure', 'Cash tied up despite accounting revenue'], ['Currency and input changes', 'Margin sensitivity and potential cash mismatches']], [2.1, 4.7])
    paragraph('FDA inspection observations require careful interpretation. A Form 483 is not a final agency determination; an Official Action Indicated classification can lead to withholding approvals and other enforcement measures. Any notice must be linked to the relevant facility, product and subsequent status before estimating its financial effect [3]. A dated regulatory event is a reason to investigate rather than an automatic numerical downgrade.')

    page('Peer groups based on operating characteristics')
    paragraph('Peer membership is an analytical input determined by business model and operating characteristics. Similar revenue sources, geography, regulatory environments, working-capital patterns and business risks make comparisons more useful. Market capitalization and the resulting credit score do not determine membership. A high-risk company and a low-risk company can be appropriate peers if their operating economics are sufficiently similar [4].')
    table(['Group', 'Companies'], [['Global Generics & Diversified (8)', "Sun Pharma; Dr. Reddy's; Cipla; Zydus Lifesciences; Lupin; Aurobindo Pharma; Glenmark Pharma; Torrent Pharma"], ['India-Focused / Branded Formulations (6)', 'Mankind Pharma; Alkem Labs; Abbott India; JB Chemicals; Ipca Labs; Ajanta Pharma'], ['API / CDMO / Contract Manufacturing (3)', "Divi's Labs; Laurus Labs; Piramal Pharma"], ['Complex / Specialty / Healthcare Platforms (3)', 'Biocon; Gland Pharma; Jubilant Pharmova']], [2.25, 4.55])
    heading('Why these comparisons are meaningful')
    paragraph('The diversified group is a benchmark for businesses exposed to multiple products and markets. The branded-formulations group emphasizes domestic commercial and franchise characteristics. The API/CDMO group focuses on manufacturing and B2B operations where investment, utilization and contract cash flows matter. The complex group brings together specialized products and healthcare platforms with distinct operating and capital requirements. The classification is intentionally broad; it does not imply identical revenue mix.')
    paragraph('Torrent remains in the diversified group under the selected project framework, while Ajanta remains in the India-focused group despite export exposure. Gland has B2B manufacturing characteristics but is placed with complex and specialty platforms because that is the intended primary comparison. Segment-level comparisons would refine these choices when detailed revenue and customer information is available.')
    heading('Benchmark construction and interpretation')
    paragraph('The dashboard calculates each issuer’s peer median from the other members of its group, excluding the issuer itself. Missing ratios are excluded from the relevant calculation. Higher-risk percentile direction depends on the ratio: higher leverage is weaker, whereas higher coverage or liquidity is stronger. The two three-company groups provide only two other-company comparators for each issuer. Their percentiles are descriptive rankings with limited statistical stability.')
    paragraph("For Laurus, Divi's and Piramal are the primary comparators. Comparing only with Sun or Mankind could confuse operating differences with relative financial weakness. Peer benchmarking supplements absolute debt-servicing and liquidity tests; it does not excuse a stressed issuer because its peers also have weak ratios.")

    page('Financial statement analysis and cash generation')
    heading('P&L analysis')
    paragraph('Revenue, EBITDA, EBIT and interest expense were organized by company and fiscal year. Year-on-year revenue growth measures expansion or contraction. EBITDA and EBIT margins show whether growth translates into operating earnings. EBIT divided by interest expense measures the earnings cushion for financing costs. Four-year CAGR between FY21 and FY25 uses four annual intervals: (FY25 revenue / FY21 revenue)^(1/4) - 1.')
    paragraph('The analysis compares changes within each issuer first, then tests its margins and risk ratios against peers. A rising revenue line with falling margins calls for investigation of pricing, mix and costs. A widening gap between EBITDA and EBIT warrants review of statement notes; the extract alone cannot establish that the difference is entirely depreciation and amortization. Tax, exceptional items and PAT are outside the available extract, so a complete net-profit quality reconciliation is not claimed.')
    heading('Balance-sheet analysis')
    paragraph('Debt/EBITDA measures gross indebtedness relative to operating earnings. Net debt/EBITDA subtracts reported cash, retaining negative values as net-cash positions. The current ratio compares current assets with current liabilities. These measures are assessed alongside receivable and inventory growth relative to revenue growth. Faster working-capital accumulation can weaken payment capacity even when sales are increasing.')
    paragraph('Liquidity depends on quality and availability, not only totals. Receivable ageing, overdue disputes, inventory obsolescence, restricted cash and near-term debt maturities are necessary follow-ups. The extract lacks payables and cost of sales, so it does not support a complete cash conversion cycle. Total current assets cannot be treated as immediately available cash.')
    heading('Cash-flow analysis')
    paragraph('Free cash flow (FCF) is defined as operating cash flow less capex. FCF/debt screens the cash surplus relative to borrowing; OCF/EBITDA screens conversion of operating earnings into cash. Negative FCF may reflect expansion or weak operations, and the reason must be established. Positive FCF supports flexibility but may also be required for tax, dividends, acquisitions or other commitments.')
    table(['Calculation', 'Credit interpretation'], [['Debt / EBITDA', 'Sensitivity of indebtedness to operating earnings'], ['EBIT / interest expense', 'Buffer against financing cost'], ['Current assets / current liabilities', 'Broad short-term balance-sheet coverage'], ['OCF - capex', 'Cash surplus after recorded capital expenditure'], ['FCF / total debt', 'Cash coverage relative to borrowing']], [2.55, 4.25])
    paragraph('A zero denominator is unavailable, not infinite. Low debt can produce extreme FCF/debt ratios; these must be interpreted with the underlying debt balance. All company monetary amounts retain source units pending confirmation of currency, scale and statement basis [1].')

    page('Internal credit ratings and early warnings')
    paragraph('The financial risk score allocates 100 possible points across six dimensions. Higher points indicate greater measured financial risk. The model is transparent screening logic rather than a statistically calibrated probability of default or an agency rating. Earnings and working-capital thresholds are project rules; model validation would require independent outcomes and a larger history [5].')
    table(['Dimension', 'Weight', 'Scoring principle'], [['Leverage', '25', 'Debt/EBITDA rises through 1x, 2x, 3x and 4x thresholds'], ['Debt servicing', '20', 'Coverage weakens through 8x, 5x, 3x and 2x thresholds'], ['Cash flow', '20', 'FCF/debt weakens through 20%, 10%, 5% and zero'], ['Liquidity', '15', 'Current ratio weakens through 1.5x, 1.2x, 1x and 0.8x'], ['Earnings resilience', '10', 'Revenue growth slows through 10%, 5%, zero and -10%'], ['Working capital', '10', 'Receivable/inventory growth increasingly exceeds sales growth']], [1.5, .65, 4.65])
    paragraph('The internal bands are VERY LOW: 0-14; LOW: 15-29; MODERATE: 30-49; HIGH: 50-69; CRITICAL: 70-100. Exact boundary inclusivity and point allocations are recorded in the project methodology. FY21 displays an 80-point observed core rather than a full rating because prior-year growth inputs are unavailable. FY22-FY25 provide 80 complete company-year scores.')
    heading('Warnings and interpretation')
    paragraph('Financial alerts flag debt/EBITDA above 3x, EBIT/interest below 3x, current ratio below 1x, negative FCF, FCF/debt below 10%, debt growth above 20%, declining revenue, receivables growing faster than revenue and a full-score increase of at least 10 points. Alerts identify review questions; multiple flags from the same underlying event should not be treated as independent evidence of default.')
    paragraph('For zero debt, the FCF/debt ratio is unavailable: nonnegative FCF receives zero cash-flow points and a cash deficit receives 20. Zero interest results in unavailable coverage and zero servicing points under the project policy. These conventions are explicit simplifying assumptions and should be checked against actual debt facilities and expenses before a credit decision.')
    heading('Separate market stress assessment')
    paragraph('Relative performance versus NIFTY Pharma, drawdown, volatility and volume contribute 35%, 30%, 20% and 15% to the separate market layer. Market deterioration can intensify review without altering the historical financial band. It may reflect equity expectations rather than inability to pay. Fresh, aligned observations are required before classifying deterioration as confirmed. There is no qualitative score in either layer.')

    page('Portfolio findings and interpretation')
    table(['FY25 internal band', 'Companies'], [['VERY LOW', '15'], ['LOW', '5'], ['MODERATE / HIGH / CRITICAL', '0']], [4.8, 2.0])
    paragraph('The low FY25 scores are a cross-sectional result, not evidence that stress never occurred. Piramal reached 75 and Jubilant 74 in FY23; Biocon reached 70 in FY23 and Laurus 70 in FY24. A time-series review identifies these adverse periods and tests whether the subsequent improvement is sustainable. Screening only the latest year would omit material historical vulnerability [1].')
    group_values = []
    for g in sorted({r['peer_group'] for r in LATEST.values()}):
        vals = [r for r in LATEST.values() if r['peer_group'] == g]
        short = {'API / CDMO / Contract Manufacturing': 'API / CDMO', 'Complex / Specialty / Healthcare Platforms': 'Complex / specialty', 'Global Generics & Diversified': 'Global / diversified', 'India-Focused / Branded Formulations': 'India / branded'}[g]
        group_values.append([short, len(vals), number(statistics.median(r['financial_risk_score'] for r in vals), 0), number(statistics.median(r['debt_ebitda'] for r in vals)) + 'x'])
    table(['Group', 'Count', 'Median score', 'Median debt / EBITDA'], group_values, [2.3, .65, 1.35, 1.9])
    paragraph('These are descriptive medians across all group members, distinct from the issuer-excluding medians used in individual benchmarking. The API/CDMO and complex groups show higher median leverage and scores in FY25. Business-model membership was fixed independently of those outcomes; the result does not explain why the groups were created.')
    heading('Review priorities at 4 October 2026')
    paragraph('The current review list contains 18 WATCH, one ENHANCED MONITORING and one CREDIT REVIEW entry. All FY25 financial observations are stale under the 540-day project policy. Mankind’s enhanced review reflects a high-severity historical financial flag; JB’s credit review reflects an entity event. Nineteen companies have a complete market score and JB lacks sufficient current history. No entry qualifies as confirmed financial-and-market deterioration because the financial evidence is not sufficiently fresh and aligned.')
    paragraph('WATCH therefore means an information or monitoring requirement, not necessarily financial weakness. A company can retain a VERY LOW historical band while requiring updated statements. Equally, a low financial score cannot clear unresolved contractual, legal-entity or regulatory questions. The review decision must state the evidence date and reason.')
    heading('Main analytical lesson')
    paragraph('Use the score to organize investigation, then explain the driver. A company with low borrowing and positive cash flow differs from one whose earnings recently recovered while debt remains significant, even if both fall within an acceptable screening band. Financial strength, exposure size, payment performance and documentation must be considered together before recommending terms or a credit limit.')

    page('Detailed case study Biocon')
    paragraph('Peer group: Complex / Specialty / Healthcare Platforms. The project compares Biocon primarily with Gland Pharma and Jubilant Pharmova. Its extract-based trajectory illustrates financial stress during growth [1, 4]. The following figures are project-input results; the external reconciliation on page 14 identifies differences from issuer disclosures, so they are not certified reported financials.')
    case_table('Biocon')
    heading('What changed and why it matters')
    paragraph('Revenue rose from 7,106 to 16,850 source units, a four-year CAGR of 24.1%. In FY23, debt increased from 5,120 to 15,200 while EBITDA was 2,528. Debt/EBITDA rose to 6.01x, coverage fell to 2.65x and FCF was -1,960. The 70-point CRITICAL score combined leverage 25, servicing 15, cash flow 20, liquidity 2 and working capital 8. Strong sales growth did not offset the financing and cash pressure.')
    paragraph('By FY25, debt fell to 11,500 and EBITDA rose to 4,450. Gross leverage improved to 2.58x, coverage to 3.17x and FCF to 1,420. The 25-point LOW score consists of leverage 12, servicing 8 and cash flow 5. Liquidity, earnings and working capital contributed zero points. Improvement reflects both a smaller debt numerator and stronger earnings; assigning a causal explanation such as an acquisition requires supporting transaction and statement notes.')
    heading('Peer assessment and credit-review recommendation')
    paragraph('FY25 leverage is higher than Gland’s 0.37x and Jubilant’s 2.07x; coverage is below their 74.44x and 3.57x respectively. Biocon’s FCF/debt of 12.3% is positive but leaves less cash coverage than Jubilant’s 20.4%. A positive cash balance reduces net leverage to 1.81x, subject to verifying that cash is unrestricted and available to the obligor.')
    paragraph('The analyst should obtain updated financials, debt maturity and covenant schedules, explain the FY23 debt step-change, and test whether recurring operating cash can fund investment and debt service. The FY25 improvement supports a more favorable historical assessment, but the remaining servicing sensitivity and stale observation date warrant evidence before an increased exposure recommendation.')

    page('Detailed case study Laurus Labs')
    paragraph("Peer group: API / CDMO / Contract Manufacturing. Laurus is benchmarked against Divi's and Piramal Pharma. Its extract-based trajectory illustrates leverage sensitivity to earnings [1, 4]. These are project-input results. Page 14 shows issuer cash-flow figures that differ materially, including a negative OCF-minus-capex outcome; the historical LOW classification remains provisional.")
    case_table('Laurus Labs')
    heading('From earnings pressure to recovery')
    paragraph('FY24 revenue fell 16.6% from 6,041 to 5,041 source units. EBITDA fell from 1,588 to 780 and the EBITDA margin narrowed from 26.3% to 15.5%. Debt rose from 2,120 to 2,380, but the earnings decline drove leverage from 1.34x to 3.05x. EBIT/interest dropped to 2.42x and FCF turned negative at -120. The CRITICAL score of 70 comprises leverage 20, servicing 15, cash flow 20, earnings 10 and working capital 5.')
    paragraph('In FY25 revenue recovered 12.7%, EBITDA rose to 1,120 and debt declined to 2,150. Leverage improved to 1.92x, coverage to 4.55x and FCF to 240. Its LOW score of 19 comprises leverage 6, servicing 8 and cash flow 5. The EBITDA margin recovered to 19.7%, still below FY23. Recovery therefore represents improvement from the stress year rather than a return to all earlier operating conditions.')
    heading('Peer assessment and decision questions')
    paragraph("Divi's shows negligible reported debt and a 32.6% FY25 EBITDA margin; its very large FCF/debt ratio reflects a tiny denominator and is unsuitable as a realistic repayment target for Laurus. Piramal shows 2.01x leverage, 4.07x coverage and 20.3% FCF/debt. Laurus is slightly stronger on gross leverage and coverage but weaker on FCF/debt at 11.2%. Ratios should be assessed together.")
    paragraph('A review should explain the FY24 margin collapse, distinguish recurring demand from temporary programs, assess customer concentration and capacity utilization, and reconcile expansion commitments with operating cash. These are investigation priorities rather than established causes. Updated statements and cash forecasts should establish whether recovered earnings and cash generation persist before relaxing payment terms or recommending a larger limit.')

    page('Company summaries Global Generics and Diversified')
    paragraph('FY25 internal classifications and ratios below come from the supplied financial extract. Each summary connects a measured result with a follow-up question. The group assignment reflects the user-defined operating framework rather than a verified identical segment mix [1, 4].')
    summary('Sun Pharma', 'FCF is positive and net leverage is -0.53x. The current ratio of 1.45x, earnings and working capital contribute two points each. Its September 2025 Halol filing described OAI and import restrictions; verify subsequent status and financial effect [8].')
    summary("Dr. Reddy's", 'Positive FCF covers 79.4% of debt and the current ratio is 2.03x. Net leverage is 0.26x. Assess investment commitments and whether strong cash generation persists; a zero score does not eliminate business or obligor risk.')
    summary('Cipla', 'Net leverage is negative and liquidity is strong at 3.71x. Five working-capital points and two earnings points explain its seven-point score. Collection and inventory trends deserve attention despite very low borrowing.')
    summary('Zydus Lifesciences', 'The current ratio is 2.88x, FCF is positive and reported cash exceeds debt. High FCF/debt mainly reflects the small debt balance. Focus review on sustainable margins, investment needs and actual cash availability.')
    summary('Lupin', 'The current ratio is 2.61x and FCF/debt is 176.2%. All six FY25 components contribute zero points. Review the durability of operating improvement and obtain current facility-specific evidence before drawing regulatory conclusions.')
    summary('Aurobindo Pharma', 'Reported net cash and 2.63x liquidity support the historical profile. Two earnings points explain the score. Verify the quality of growth and how capex, working capital and site dependency affect future cash.')
    summary('Glenmark Pharma', 'Positive FCF and reported net cash support the financial result. The FDA issued a warning letter dated 11 July 2025 [9]. That dated finding warrants scope and resolution checks; it does not establish an unresolved issue in October 2026.')
    summary('Torrent Pharma', 'Liquidity is 1.84x and FCF/debt 87.1%; two working-capital points explain the score. The FY25 extract predates the July 2026 JB amalgamation. Evaluate updated group debt, combined operations and legal exposure separately [6].')

    page('Company summaries Branded and specialized businesses')
    heading('India-Focused / Branded Formulations')
    summary('Mankind Pharma', 'Leverage contributes 12 points and working capital eight. Coverage and positive FCF are favorable, but debt growth warrants enhanced monitoring. Explain funding changes and reconcile receivables and inventory with sales.')
    summary('Alkem Labs', 'Reported net cash and 2.49x liquidity support flexibility. Earnings and working capital contribute two points each. Investigate cash conversion rather than interpreting the four-point score alone.')
    summary('Abbott India', 'Reported debt is zero, the current ratio is 3.26x and FCF is positive. FCF/debt is unavailable, not infinite. Two earnings points explain its score; verify unrestricted cash and contractual obligor identity.')
    summary('JB Chemicals', 'FCF is positive and liquidity is 2.79x in FY25. These historical figures do not represent successor financials after amalgamation. Reconcile outstanding obligations, entity mapping and successor documentation [6].')
    summary('Ipca Labs', 'Leverage contributes all six points. Liquidity is 2.56x and FCF/debt 47.0%. Review borrowing purpose and maturities while assessing whether cash generation supports ongoing investment.')
    summary('Ajanta Pharma', 'The debt balance is small and reported cash exceeds debt. Two working-capital points explain its score. Large coverage ratios reflect small interest and debt denominators; inspect collection quality and geographic mix.')
    heading('Remaining API and complex businesses')
    summary("Divi's Labs", 'Debt is just one source unit; leverage rounds to 0.00x and liquidity is 7.16x. FCF/debt is exceptionally large because debt is tiny. Positive FCF is the more useful indicator; assess future investment and customer dependency.')
    summary('Piramal Pharma', 'Leverage, servicing and working capital contribute 12, eight and two points. FCF/debt is 20.3%. FY23 reached CRITICAL at 75; confirm the durability of recovery and debt maturity coverage.')
    summary('Gland Pharma', 'Reported net cash and 3.58x liquidity support the profile. Two working-capital points explain the score. Review receivable quality, inventory and investment obligations within its specialized business model.')
    summary('Jubilant Pharmova', 'Leverage, servicing and working capital contribute 12, eight and two points. FY23 reached CRITICAL at 74. Positive FY25 FCF improves flexibility, while coverage remains materially below the strongest peers.')

    page('Counterparty assessment and credit reviews')
    heading('Identify who is expected to pay')
    paragraph('Counterparty analysis starts with the legal obligor rather than the brand or listed group. Establish the contracting entity, ownership, operating role, jurisdiction, related entities and any enforceable guarantee. A parent’s strong financial position does not automatically protect a subsidiary obligation. Record the transaction purpose, currency, payment term and exposure type before interpreting the financial score.')
    heading('Evaluate capacity and willingness to meet obligations')
    paragraph('Capacity is assessed through recurring earnings, operating cash, leverage, interest servicing, near-term liquidity and funding access. Payment performance, disputed invoices, covenant compliance and management responses supplement the financial picture. Concentration must be considered at both issuer and economic-group level so that multiple subsidiaries do not conceal dependence on one source of repayment.')
    paragraph('The Basel Committee’s 2025 principles emphasize an appropriate credit-risk environment, sound credit granting, effective administration and monitoring, and adequate controls [7]. These principles inform the proposed governance approach; they are not presented as binding rules for this student project or as a substitute for the institution’s own credit policy.')
    heading('Review process')
    paragraph('A scheduled review updates statements, ratios, score drivers and exposure records. An event-driven review investigates a material earnings decline, negative FCF, leverage or coverage stress, debt acceleration, overdue payments, market stress, regulatory developments or changes in legal identity. Analyst judgment establishes whether the trigger reflects temporary volatility, structural weakening or an information gap.')
    table(['Review stage', 'Required outcome'], [['Evidence collection', 'Dated statements, legal identity and exposure reconciliation'], ['Analysis', 'Ratio trajectory, score drivers, peer differences and cash forecast'], ['Challenge', 'Downside assumptions, concentration and protection assessment'], ['Recommendation', 'Retain or amend terms, limit proposal and review frequency'], ['Decision and follow-through', 'Authorized approval, conditions, action owner and due date']], [1.9, 4.9])
    heading('Credit memo and recommendation')
    paragraph('The memo should state the obligor and statement scope, summarize current capacity, explain historical stress, identify business risks, reconcile exposure and record security or guarantees. It should give a reasoned recommendation with conditions and an accountable approval route. The score is supporting evidence, not the decision itself. A favorable historical band can coexist with a recommendation to defer increased exposure until financials are refreshed.')
    paragraph('JB illustrates the need for entity-aware review: Torrent’s filing states that amalgamation became effective on 8 July 2026 [6]. Historical issuer records should be retained while successor obligations and combined exposure are checked. A merger event does not justify replacing historical ratios with assumed successor ratios or automatically duplicating credit limits.')

    page('Credit-limit monitoring and exposure controls')
    paragraph('This project specifies a proposed monitoring workflow. Actual approved limits, transaction exposures and overdue balances have not been supplied, so the current dashboard has not measured real limit utilization or breaches. Financial screening determines where investigation is needed; limit setting also requires transaction size, tenor, payment evidence, protection and authorized policy.')
    heading('Define and reconcile exposure')
    paragraph('For trade credit, start with unpaid invoices and policy-defined shipped-but-unbilled obligations. For lending, use outstanding principal, accrued obligations and policy-defined commitments. Keep drawn and undrawn components distinct and avoid counting the same obligation twice. Aggregate by legal entity and economic group, using consistent currency conversion and a dated snapshot.')
    table(['Measure', 'Definition'], [['Utilization', 'Exposure / approved limit x 100'], ['Headroom', 'Approved limit - exposure'], ['Breach amount', 'Maximum of exposure - approved limit and zero'], ['Overdue exposure', 'Unpaid amount beyond its contractual due date'], ['Review status', 'Validity of approval, expiry and next review date']], [1.6, 5.2])
    paragraph('Illustration only: an approved limit of INR 10 million and exposure of INR 8 million imply 80% utilization and INR 2 million headroom. Exposure of INR 11 million would imply 110% utilization and an INR 1 million breach. These numbers are hypothetical, not limits approved for a studied company. Missing or zero limits are exceptions requiring review; they do not imply healthy utilization.')
    heading('Monitoring cadence and escalation')
    paragraph('A proposed operating approach reconciles exposure at each agreed snapshot, reviews overdue balances and upcoming expiries, and escalates breaches to the authorized credit owner. An early-warning level such as 80% could be adopted only after policy approval; 100% identifies the arithmetic limit boundary. Approvals, temporary exceptions and override expiry should be recorded and reviewed. Analysts recommend actions; authorized decision-makers approve terms and limit changes.')
    paragraph('A breach review checks whether the exposure is accurate, whether an exception exists, whether further transactions should be restricted under policy, and whether collections or an approved limit revision is required. Deteriorating financials can justify a reassessment even below the limit. Good financials do not excuse an unauthorized breach.')
    heading('Data needed to activate the framework')
    paragraph('Required fields are legal-entity and group identifiers, exposure type, currency, dated outstanding and committed amounts, approved limit, approver, effective and expiry dates, payment terms, overdue ageing, guarantees or collateral, eligible value and haircuts, review date and exception status. Gross and eligible net exposure should remain distinguishable. Collateral is deducted only under a documented valuation and enforceability policy.')

    page('Benchmark comparisons and external credit context')
    heading('Observed peer comparisons')
    paragraph('FY25 leverage is compared with the median of the other companies in each business-model group, excluding the issuer. A positive gap means higher gross leverage than the comparators; it is not a default probability. These are actual calculations on the project extract, rather than generic target ratios [1].')
    examples = []
    for name in ['Biocon', 'Laurus Labs', 'Mankind Pharma']:
        r = LATEST[name]
        examples.append([name, number(r['debt_ebitda'])+'x', number(r['debt_ebitda_peer_median'])+'x', f"{r['debt_ebitda']-r['debt_ebitda_peer_median']:+.2f}x"])
    table(['Company', 'Debt / EBITDA', 'Other-peer median', 'Gap'], examples, [1.65, 1.5, 1.8, 1.25])
    heading('Actual sector-index comparison')
    paragraph('NSE Indices identifies NIFTY Pharma as a pharmaceutical-sector performance benchmark [10]. The cached market snapshot compares company and index returns over identical endpoints: 25 May to 1 October 2026, covering 90 company trading-session intervals. Nineteen companies have complete market scores. The following figures are recomputed from the cached price data, not returns independently downloaded from NSE.')
    examples = []
    markets = {r['company_id']: r for r in DATA['v2']['market']}
    for name in ['Sun Pharma', 'Biocon', 'Laurus Labs']:
        r = markets[name]
        examples.append([name, f"{r['return_90']:.2%}", f"{r['benchmark_return_90']:.2%}", f"{r['relative_return_90']*100:+.2f} pp"])
    table(['Company', 'Company return', 'NIFTY Pharma', 'Relative return'], examples, [1.65, 1.5, 1.8, 1.25])
    paragraph('Equity adjusted closes include dividend/split adjustments, while the cached index is a price index. This return-convention mismatch limits precision; a total-return-index series would improve comparability. Index membership is separate from the four business-model peer groups. Market underperformance supports investigation, not a financial-rating downgrade.')
    heading('External rating comparison')
    paragraph('CARE reaffirmed Laurus long-term bank facilities at CARE AA; Stable on 1 July 2025, revising the outlook from Negative [11]. The rationale describes improved FY25 profitability and debt coverage, directionally consistent with the project’s recovery assessment. This is a dated contextual check, not an agency-equivalent validation of the 19/100 LOW score or evidence of the current October 2026 rating.')
    paragraph('CARE also identifies sustained total debt/PBILDT above 2.75x as a negative sensitivity. The project’s leverage alert uses debt/EBITDA above 3x. These thresholds use different analytical definitions and purposes; they must not be treated as interchangeable. One issuer comparison cannot establish rating accuracy across the portfolio.')

    page('Validation findings and source reconciliation')
    heading('Comparison with official FY25 disclosures')
    table(['Issuer and metric', 'Project extract', 'Issuer disclosure'], [
        ['Biocon total revenue', '16,850', '16,470'],
        ['Biocon EBITDA', '4,450', '4,374'],
        ['Laurus revenue', '5,680', '5,554'],
        ['Laurus EBITDA', '1,120', '1,115'],
        ['Laurus operating cash flow', '890', '602'],
        ['Laurus capex', '650', '659'],
    ], [3.25, 1.55, 1.95])
    paragraph('Disclosed amounts are INR crore; project monetary units and statement basis remain unconfirmed. Biocon figures are from its 8 May 2025 results release [12]. Laurus figures are from its 24 April 2025 presentation, printed page 17 [13]. These comparisons identify unreconciled numeric differences, not a like-for-like source validation pass.')
    paragraph('Biocon also discloses operating revenue of 15,262, distinct from total revenue of 16,470. Its revenue definition must be reconciled before growth and margin comparisons. Laurus discloses net debt/EBITDA of 2.3x, compared with 1.76x in the project. More materially, disclosed OCF minus capex is -57, whereas the extract gives +240. The different cash-flow sign could affect the financial band; no hybrid score is calculated by replacing only selected fields.')
    heading('Calculation and consistency checks completed')
    audit = json.loads((OUT / 'result_validation.json').read_text())
    c = audit['counts']
    paragraph(f"Independent recomputation from project CSV inputs passed {c['financial_metric_checks']} financial-metric comparisons across 100 records, {c['score_reconciliations']} component-sum reconciliations, {c['peer_median_checks']} FY25 issuer-excluding leverage medians and {c['market_endpoint_checks']} return/benchmark/relative-return comparisons across 19 issuers. That is {sum(c.values())} numerical comparisons. The checker does not call the production scoring or market functions.")
    paragraph('All 28 regression tests passed. They cover independent formula examples, score and alert boundaries, FY21 unavailable growth, zero denominators, peer exclusion/ties, market windows and benchmark endpoints, freshness, entity transitions and full-pipeline behavior. The 19 financial input checks pass structural and internal-consistency rules; they do not authenticate the monetary fields against audited statements [14].')
    heading('What validation does and does not establish')
    paragraph('Arithmetic verification passes against the supplied inputs. External source reconciliation remains open and issuer results are provisional. Default-prediction accuracy, agency-grade agreement, false-positive rates and out-of-sample performance have not been established: the study lacks matched outcomes and sufficient calibration history. Before operational use, reconcile the full statements, rerun all affected metrics and scores, and then validate against dated rating/default outcomes on a separate period.')

    page('Evidence standards and report references')
    heading('Interpretation boundaries')
    paragraph('The financial extract contains 100 observations across 20 issuers and five fiscal years. Currency, monetary scale, consolidated or standalone scope and statement provenance require reconciliation before monetary amounts are used to size credit. The analysis is a standardized ratio study, not an audit or reconstruction of every statement line. FY25 figures are historical relative to the October 2026 review date.')
    paragraph('Missing statement detail includes PAT, exceptional items, equity, cost of sales, payables, debt maturity schedules, contingent liabilities and complete financing/investing cash flows. These gaps limit earnings-quality, solvency, cash-conversion-cycle and refinancing conclusions. The internal rating is uncalibrated screening logic and does not measure default probability, recovery or an external agency grade.')
    paragraph('Business-model groups provide useful comparisons but retain issuer-specific differences. Two groups have only three members. Narrative company material is retained for research with source dates and locations, without a qualitative score. User narratives require corroboration; older regulatory findings do not establish current status. Market signals and annual financials have different observation periods and must not be combined as if contemporaneous.')
    heading('References and source hierarchy')
    paragraph('[1] Project financial source: pharma_financials_5yr.csv; calculated metrics and scores in outputs/dashboard_data.json. FY21-FY25. Monetary fields retain source units. Company ratios, score components and historical findings in this report are derived from this source.')
    paragraph('[2] Government of India, Department of Pharmaceuticals, Annual Report 2025-26, printed page 3 (PDF page 10). Sector figures relate to FY2024-25. https://pharma-dept.gov.in/sites/default/files/Annual%20Report%202025-26_0.pdf')
    paragraph('[3] US Food and Drug Administration, Pharmaceutical Inspections and Compliance. Accessed 4 October 2026. https://www.fda.gov/drugs/guidance-compliance-regulatory-information/pharmaceutical-inspections-and-compliance')
    paragraph('[4] Project peer-group rationale supplied by the project owner; assignment register config/peer_groups.csv and discussion docs/PEER_GROUPS.md. Membership is based on operating characteristics, not financial scores or market capitalization.')
    paragraph('[5] Project methodology: docs/METHODOLOGY.md and score configuration. Review-priority and market rules: config/monitoring.json. Dataset and narrative provenance: README.md and data/README.md.')
    paragraph('[6] Torrent Pharmaceuticals, Intimation of Effective Date of Scheme of Amalgamation, 8 July 2026. https://www.torrentpharma.com/docs/02_Intimation_Of_Effective_Date_Of_Scheme_Of_Amalgamation_b3199d2259.pdf')
    paragraph('[7] Basel Committee on Banking Supervision, Principles for the Management of Credit Risk, 30 April 2025. https://www.bis.org/publications/202504-guidelines-principles-management-credit-risk')
    paragraph('[8] Sun Pharmaceutical Industries, Halol status filing, 9 September 2025. https://sunpharma.com/wp-content/uploads/2025/09/SEIntimationHalolUpdate.pdf')
    paragraph('[9] FDA, Glenmark Pharmaceuticals warning letter, 11 July 2025. https://www.fda.gov/inspections-compliance-enforcement-and-criminal-investigations/warning-letters/glenmark-pharmaceuticals-limited-708270-07112025')
    paragraph('[10] NSE Indices, NIFTY Pharma index description. Accessed 4 October 2026. https://www.niftyindices.com/indices/equity/sectoral-indices/nifty-pharma')
    paragraph('[11] CARE Ratings, Laurus Labs Limited rating rationale, 1 July 2025, page 1. https://www.careratings.com/upload/CompanyFiles/PR/202507120700_Laurus_Labs_Limited.pdf')
    paragraph('[12] Biocon, FY25 consolidated financial results release, 8 May 2025. https://www.biocon.com/biocon-q4fy25-revenue/')
    paragraph('[13] Laurus Labs, FY25 results presentation, 24 April 2025, printed page 17. https://www.lauruslabs.com/Investors/PDF/Q4/InvPresentation24042025.pdf')
    paragraph('[14] Validation evidence: reports/result_validation.json; independent checker scripts/validate_report_results.py; regression suite tests/test_crm.py and tests/test_v2.py. Tests passed on 4 October 2026. Source discrepancies are preserved in reports/external_benchmarks.json.')
    OUT.mkdir(exist_ok=True)
    DOC.save(OUT / 'Indian_Pharma_Counterparty_Credit_Report.docx')
    (OUT / 'Indian_Pharma_Counterparty_Credit_Report.md').write_text('\n'.join(MD))
    print('Created DOCX and Markdown report')


if __name__ == '__main__':
    build()
